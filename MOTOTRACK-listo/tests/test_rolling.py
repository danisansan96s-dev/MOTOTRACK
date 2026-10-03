"""Actualización rolling con los Excel reales de 208/209 y copias sintéticas posteriores."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from streamlit.delta_generator import DeltaGenerator
from streamlit.testing.v1 import AppTest

from data_processing import load_data, analyze, PRODUCTS
from forecasting import evaluate, forecast, predict, MODELS

ROOT = Path(__file__).resolve().parents[1]
EXPECTED_SERIES = {(region, product) for region, products in PRODUCTS.items() for product in products}
REGION_COLUMNS = {'Norte': [1, 2, 3, 4], 'Centro': [6, 7, 8, 9], 'Sur': [11, 12, 13]}


def workbook_for(period_count, region_counts=None):
    """Las extensiones de 210 en adelante son exclusivamente datos de prueba."""
    source = ROOT / ('moto-track.xlsx' if period_count == 208 else 'moto-track-209.xlsx')
    book = load_workbook(source)
    sheet = book['Demand']
    for region, columns in REGION_COLUMNS.items():
        count = (region_counts or {}).get(region, period_count)
        for period in range(210, count + 1):
            sheet.cell(period + 2, columns[0], period)
            for column in columns[1:]:
                sheet.cell(period + 2, column, sheet.cell(211, column).value + (period - 209) % 3)
        for row in range(count + 3, sheet.max_row + 1):
            for column in columns:
                sheet.cell(row, column).value = None
    content = BytesIO()
    book.save(content)
    book.close()
    return content.getvalue()


class RollingTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cases = {}
        for count in [208, 209, 210]:
            content = workbook_for(count)
            data, _ = load_data(content)
            comparison, selected, validation = evaluate(data)
            detail, consolidated, _ = forecast(data, selected)
            cls.cases[count] = (content, data, comparison, selected, validation, detail, consolidated)

    def test_rolling_ranges_and_models(self):
        for count, (_, data, comparison, selected, validation, detail, total) in self.cases.items():
            with self.subTest(periods=count):
                self.assertEqual(len(data), count * 8)
                self.assertEqual(len(comparison), len(MODELS) * 8)
                self.assertEqual(set(data[['Regional', 'Producto']].itertuples(index=False, name=None)), EXPECTED_SERIES)
                self.assertEqual(len(selected), 8)
                self.assertEqual(total.Periodo.tolist(), list(range(count + 1, 227)))
                self.assertEqual(len(detail), (226 - count) * 8)
                self.assertTrue((detail['Pronóstico'] >= 0).all())
                self.assertEqual(set(validation.Periodo), set(range(count - 51, count + 1)))
                self.assertFalse(((detail.Regional == 'Sur') & (detail.Producto == 'Tractor')).any())
                self.assertTrue((analyze(data).Observaciones == count).all())
                with self.assertRaisesRegex(ValueError, '226'):
                    forecast(data, selected, 227 - count)

    def test_new_observation_used_to_refit(self):
        _, data, _, selected, _, _, _ = self.cases[209]
        with patch('forecasting.predict', wraps=predict) as fitting:
            forecast(data, selected)
        self.assertEqual(len(fitting.call_args_list), 8)
        for call, row in zip(fitting.call_args_list, selected.itertuples()):
            expected = data[(data.Regional == row.Regional) & (data.Producto == row.Producto)].sort_values('Periodo').Demanda.to_numpy()
            np.testing.assert_array_equal(call.args[0], expected)
            self.assertEqual(len(call.args[0]), 209)
            self.assertEqual(call.args[0][-1], expected[-1])
            self.assertEqual(call.args[2], 17)
        # Naive debe reflejar exactamente el último real, no el periodo 208.
        naive = selected.copy()
        naive.Modelo = 'Naive'
        detail, _, _ = forecast(data, naive, 1)
        expected = data[data.Periodo == 209].sort_values(['Regional', 'Producto']).Demanda.to_numpy()
        np.testing.assert_array_equal(detail.sort_values(['Regional', 'Producto'])['Pronóstico'], expected)

    def test_dynamic_validation_training_window(self):
        for count in [209, 210]:
            data = self.cases[count][1]
            # Modelo determinista para inspeccionar cada corte temporal sin
            # repetir las optimizaciones verificadas en test_rolling_ranges.
            def tracked(y, model, horizon):
                self.assertEqual(len(y), count - 52)
                self.assertEqual(horizon, 52)
                return np.repeat(y[-1], horizon), ''
            with patch('forecasting.predict', side_effect=tracked) as fitting:
                _, _, validation = evaluate(data)
            self.assertEqual(fitting.call_count, 80)
            self.assertEqual(set(validation.Periodo), set(range(count - 51, count + 1)))

    def test_region_alignment_and_minimum(self):
        content = workbook_for(209, {'Norte': 209, 'Centro': 208, 'Sur': 208})
        with self.assertRaisesRegex(ValueError, r'Regionales desalineadas.*Norte: 209.*Centro: 208.*Sur: 208'):
            load_data(content)
        with self.assertRaisesRegex(ValueError, 'mínimo de 208'):
            load_data(workbook_for(207))

    def test_last_playable_period_and_completed_game(self):
        data, _ = load_data(workbook_for(225))
        selected = self.cases[209][3].copy()
        selected.Modelo = 'Naive'
        detail, total, _ = forecast(data, selected)
        self.assertEqual(total.Periodo.tolist(), [226])
        self.assertEqual(len(detail), 8)
        with self.assertRaises(ValueError):
            forecast(data, selected, 2)
        completed, _ = load_data(workbook_for(226))
        with self.assertRaisesRegex(ValueError, 'finalizó'):
            forecast(completed, selected)

    def test_upload_recomputes_horizon_and_results(self):
        # La fuente cambia dentro de la misma sesión, como ocurre al volver a
        # exportar el juego. El cache depende del contenido de cada Excel.
        source = SimpleNamespace(getvalue=lambda: self.cases[208][0])
        with patch.object(DeltaGenerator, 'file_uploader', side_effect=lambda *args, **kwargs:
                          None if kwargs.get('key') == 'costs_upload' else source):
            app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180).run()
            for count in [208, 209, 210]:
                source.getvalue = lambda n=count: self.cases[n][0]
                app.run()
                self.assertFalse(app.exception, str(app.exception))
                self.assertFalse(app.error, str(app.error))
                self.assertEqual(app.metric[0].value, str(count))
                self.assertEqual(app.slider[0].max, 226 - count)
                self.assertEqual(app.slider[0].value, 226 - count)
                actual = app.tabs[2].dataframe[4].value
                pd.testing.assert_frame_equal(actual, self.cases[count][5])
                self.assertTrue(any(f'próximo periodo a pronosticar: {count + 1}' in c.value for c in app.caption))
            source.getvalue = lambda: workbook_for(225)
            app.run()
            self.assertFalse(app.exception, str(app.exception))
            self.assertFalse(app.error, str(app.error))
            self.assertEqual(app.slider[0].max, 1)
            self.assertEqual(app.slider[0].value, 1)
            self.assertTrue(app.slider[0].disabled)
            self.assertEqual(set(app.tabs[2].dataframe[4].value.Periodo), {226})
            source.getvalue = lambda: workbook_for(226)
            app.run()
            self.assertFalse(app.exception, str(app.exception))
            self.assertFalse(app.error, str(app.error))
            self.assertFalse(app.slider)
            self.assertTrue(any('finalizó' in item.value for item in app.info))


if __name__ == '__main__':
    unittest.main()
