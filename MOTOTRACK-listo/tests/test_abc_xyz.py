from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from abc_xyz import (load_costs, persist_costs, costs_path, classify, conclusions, MAIN_COLUMNS, SCORE_FORMULA)
from data_processing import load_data
from forecasting import evaluate, metrics
from export import export_abc_excel
from visualizations import pareto_abc, abc_xyz_matrix
from test_rolling import workbook_for

ROOT = Path(__file__).resolve().parents[1]


def cost_content(frame):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        frame.to_excel(writer, sheet_name='Costos', index=False)
    return buffer.getvalue()


class AbcXyzTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cost_bytes = (ROOT / 'Costos_MotoTrak_Actividad.xlsx').read_bytes()
        cls.costs = load_costs(cls.cost_bytes)
        cls.data, _ = load_data((ROOT / 'moto-track-209.xlsx').read_bytes())
        _, cls.selected, cls.validation = evaluate(cls.data)
        cls.table, cls.total = classify(cls.data, cls.costs, cls.selected, cls.validation)

    def test_real_costs_and_direct_unit_profit(self):
        costs = self.costs.set_index('Producto')
        self.assertEqual(costs.loc['Moto', 'MD'], 3947000)
        self.assertEqual(costs.loc['Moto', 'MOD'], 1642500)
        self.assertEqual(costs.loc['Moto', 'Costo fabricación'], 5589500)
        self.assertEqual(costs.loc['Cuatrimoto', 'Utilidad unitaria'], 4673500)
        changed = self.costs.copy()
        changed.loc[changed.Producto == 'Moto', 'Utilidad unitaria'] = 12345
        loaded = load_costs(cost_content(changed))
        self.assertEqual(loaded.loc[loaded.Producto == 'Moto', 'Utilidad unitaria'].iloc[0], 12345)
        fallback = load_costs(cost_content(changed.drop(columns='Utilidad unitaria')))
        np.testing.assert_array_equal(fallback['Utilidad unitaria'], fallback['Precio venta'] - fallback['Costo fabricación'])

    def test_persistent_valid_costs_and_invalid_upload(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            destination = root / 'data' / 'costos_actuales.xlsx'
            self.assertEqual(costs_path(root).name, 'Costos_MotoTrak_Actividad.xlsx')
            persist_costs(self.cost_bytes, destination)
            self.assertEqual(costs_path(root), destination)
            self.assertEqual(destination.read_bytes(), self.cost_bytes)
            replacement = self.costs.copy()
            replacement.loc[0, 'MD'] += 1
            content = cost_content(replacement)
            persist_costs(content, destination)
            self.assertEqual(destination.read_bytes(), content)
            with self.assertRaises(ValueError):
                persist_costs(cost_content(replacement.drop(columns='MD')), destination)
            self.assertEqual(destination.read_bytes(), content)

    def test_exact_eight_skus_window_profit_and_score(self):
        table = self.table
        self.assertEqual(table.columns.tolist(), MAIN_COLUMNS)
        self.assertEqual(len(table), 8)
        self.assertEqual(table.SKU.nunique(), 8)
        self.assertFalse(((table.Regional == 'Sur') & (table.Producto == 'Tractor')).any())
        self.assertTrue(table['Periodo inicial ventana'].eq(158).all())
        self.assertTrue(table['Periodo final ventana'].eq(209).all())
        window = self.data[self.data.Periodo.between(158, 209)]
        self.assertTrue(window.groupby(['Regional', 'Producto']).size().eq(52).all())
        for row in table.to_dict('records'):
            real = window[(window.Regional == row['Regional']) & (window.Producto == row['Producto'])].Demanda.sum()
            self.assertEqual(row['Demanda 52 semanas'], real)
            self.assertEqual(row['Utilidad total 52 semanas'], real * row['Utilidad unitaria'])
            selected = self.selected[(self.selected.Regional == row['Regional']) & (self.selected.Producto == row['Producto'])].iloc[0]
            self.assertAlmostEqual(row['Score%'], selected.WMAPE * 100)
            self.assertEqual(row['RMSE'], selected.RMSE)
            self.assertEqual(row['Modelo seleccionado'], selected.Modelo)
        self.assertAlmostEqual(table['Participación utilidad %'].sum(), 100)
        self.assertAlmostEqual(table['Participación acumulada %'].iloc[-1], 100)

    def test_abc_uses_profit_and_crossing_sku(self):
        costs = self.costs.copy()
        costs.loc[costs.Producto == 'Moto', 'Utilidad unitaria'] = 1
        table, _ = classify(self.data, costs, self.selected, self.validation)
        self.assertTrue(table.iloc[:5].Producto.ne('Moto').all())
        previous = table['Participación acumulada %'].shift(fill_value=0)
        expected = np.select([previous < 80, previous < 95], ['A', 'B'], default='C')
        np.testing.assert_array_equal(table.ABC, expected)
        small_limits, _ = classify(self.data, self.costs, self.selected, self.validation, a_limit=1, b_limit=2)
        self.assertEqual(small_limits.ABC.tolist(), ['A'] + ['C'] * 7)
        changed, _ = classify(self.data, self.costs, self.selected, self.validation, a_limit=30, b_limit=60)
        self.assertFalse(changed.ABC.equals(self.table.ABC))

    def test_xyz_uses_forecast_error_and_thresholds(self):
        table, _ = classify(self.data, self.costs, self.selected, self.validation, x_limit=0, y_limit=1)
        self.assertTrue(table.XYZ.eq('Z').all())
        table, _ = classify(self.data, self.costs, self.selected, self.validation, x_limit=100, y_limit=101)
        self.assertTrue(table.XYZ.eq('X').all())
        # Aumentar la dispersión fuera de la ventana no afecta XYZ; solo se
        # emplean las métricas validadas sobre las últimas 52 semanas.
        changed = self.data.copy()
        changed.loc[changed.Periodo < 158, 'Demanda'] *= 50
        table, _ = classify(changed, self.costs, self.selected, self.validation)
        pd.testing.assert_series_equal(table.XYZ, self.table.XYZ)
        np.testing.assert_array_equal(table['Score%'], self.table['Score%'])
        mismatched = self.validation.copy()
        mismatched.loc[mismatched.Periodo == 209, 'Real'] += 1
        with self.assertRaisesRegex(ValueError, 'reales'):
            classify(self.data, self.costs, self.selected, mismatched)

    def test_dynamic_windows_214_and_215(self):
        for period, first in [(214, 163), (215, 164)]:
            data, _ = load_data(workbook_for(period))
            _, selected, validation = evaluate(data)
            table, _ = classify(data, self.costs, selected, validation)
            self.assertTrue(table['Periodo inicial ventana'].eq(first).all())
            self.assertTrue(table['Periodo final ventana'].eq(period).all())
            expected = data[data.Periodo.between(first, period)].groupby(['Regional', 'Producto']).Demanda.sum()
            for row in table.to_dict('records'):
                self.assertEqual(row['Demanda 52 semanas'], expected.loc[(row['Regional'], row['Producto'])])

    def test_consolidation_and_aggregate_errors(self):
        for row in self.total.to_dict('records'):
            product = row['Producto']
            sku = self.table[self.table.Producto == product]
            self.assertEqual(row['Demanda 52 semanas'], sku['Demanda 52 semanas'].sum())
            self.assertEqual(row['Utilidad total 52 semanas'], sku['Utilidad total 52 semanas'].sum())
            validation = self.validation.merge(self.selected[['Regional', 'Producto', 'Modelo']], on=['Regional', 'Producto', 'Modelo'])
            totals = validation[validation.Producto == product].groupby('Periodo')[['Real', 'Pronóstico']].sum()
            expected = metrics(totals.Real, totals['Pronóstico'])
            self.assertAlmostEqual(row['Score% consolidado'], expected['WMAPE'] * 100)
            self.assertAlmostEqual(row['RMSE consolidado'], expected['RMSE'])
        self.assertAlmostEqual(self.total['Porcentaje demanda total %'].sum(), 100)
        self.assertAlmostEqual(self.total['Participación utilidad %'].sum(), 100)

    def test_excel_and_graphs(self):
        draft = conclusions(self.table, 80, 95, 15, 30)
        payload = export_abc_excel(self.costs, self.table, self.total, draft, {'A': 80, 'B': 95, 'X': 15, 'Y': 30})
        book = load_workbook(BytesIO(payload))
        self.assertTrue({'Costos', 'Clasif ABC_XYZ', 'Conclusiones', 'Consolidado', 'Pronosticos'}.issubset(book.sheetnames))
        self.assertEqual(book['Clasif ABC_XYZ'].max_row, 11)
        for sheet in book:
            self.assertFalse(sheet._charts)
            self.assertFalse(sheet._images)
            self.assertEqual(sheet.freeze_panes, 'A4')
            self.assertTrue(sheet.auto_filter.ref)
        score_column = MAIN_COLUMNS.index('Score%') + 1
        cell = book['Clasif ABC_XYZ'].cell(4, score_column)
        self.assertAlmostEqual(cell.value, self.table['Score%'].iloc[0])
        self.assertEqual(cell.number_format, '0.00" %"')
        models = pd.read_excel(BytesIO(payload), sheet_name='Pronosticos', header=2)
        self.assertTrue(models['Fórmula Score%'].eq(SCORE_FORMULA).all())
        costs = pd.read_excel(BytesIO(payload), sheet_name='Costos', header=2)
        pd.testing.assert_frame_equal(costs, self.costs)
        self.assertEqual(len(pareto_abc(self.table, 80, 95).data), 2)
        self.assertEqual(sum(sum(row) for row in abc_xyz_matrix(self.table).data[0].z), 8)

    def test_ui_parameter_updates_keep_forecaster(self):
        app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=180).run()
        self.assertFalse(app.exception, str(app.exception))
        self.assertFalse(app.error, str(app.error))
        self.assertEqual([tab.label for tab in app.tabs], ['Histórico', 'Validación de modelos', 'Pronóstico', 'ABC-XYZ', 'Consolidado', 'Conclusiones'])
        original = app.tabs[2].dataframe[4].value.copy()
        first = app.tabs[3].dataframe[0].value.copy()
        app.number_input(key='abc_limit_a').set_value(30.)
        app.number_input(key='abc_limit_b').set_value(60.)
        app.run()
        self.assertFalse(app.exception, str(app.exception))
        self.assertFalse(app.error, str(app.error))
        self.assertFalse(app.tabs[3].dataframe[0].value.ABC.equals(first.ABC))
        app.number_input(key='xyz_limit_x').set_value(100.)
        app.number_input(key='xyz_limit_y').set_value(101.)
        app.run()
        self.assertTrue(app.tabs[3].dataframe[0].value.XYZ.eq('X').all())
        pd.testing.assert_frame_equal(app.tabs[2].dataframe[4].value, original)
        self.assertEqual(len(app.get('download_button')), 2)
        app.number_input(key='abc_limit_a').set_value(70.)
        app.number_input(key='abc_limit_b').set_value(60.)
        app.run()
        self.assertTrue(app.error)
        pd.testing.assert_frame_equal(app.tabs[2].dataframe[4].value, original)


if __name__ == '__main__':
    unittest.main()
