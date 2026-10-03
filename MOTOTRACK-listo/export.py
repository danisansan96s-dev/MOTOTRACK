from io import BytesIO
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill


def export_excel(detail, consolidated, selected, comparison):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for name, frame in [('Pronostico_Detallado',detail),('Consolidado',consolidated),('Modelos',selected),('Comparacion_Modelos',comparison)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]; sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for k, col in enumerate(frame.columns, 1):
                sheet.column_dimensions[get_column_letter(k)].width = min(55, max(16,len(str(col))+2))
                if col in ['WMAPE','MAPE']:
                    for cell in list(sheet.columns)[k-1][1:]: cell.number_format = '0.00%'
    return buffer.getvalue()


def export_abc_excel(costs, classification, consolidated, conclusions, parameters):
    """Reporte tabular ABC-XYZ: porcentajes en puntos porcentuales, sin gráficos."""
    from abc_xyz import SCORE_FORMULA, ABC_RULE
    buffer = BytesIO()
    models = classification[['SKU', 'Regional', 'Producto', 'Periodo inicial ventana', 'Periodo final ventana',
                             'Modelo seleccionado', 'Score%', 'RMSE']].copy()
    models['Fórmula Score%'] = SCORE_FORMULA
    metadata = pd.DataFrame([
        ('Score%', SCORE_FORMULA), ('RMSE', 'Unidades de demanda; no es porcentaje.'),
        ('ABC', 'Demanda de las últimas 52 semanas × utilidad unitaria del archivo.'),
        ('Cruce de umbral ABC', ABC_RULE),
        ('Consolidado', 'Error de la suma de pronósticos de los modelos elegidos frente a la suma de los reales; no se promedian RMSE.'),
        ('Ventana', f'{classification["Periodo inicial ventana"].iloc[0]}–{classification["Periodo final ventana"].iloc[0]}, 52 semanas por SKU.'),
        ('Costos', 'Se conservan MD, MOD, costo, precio y utilidad del archivo utilizado. Solo si falta la columna Utilidad unitaria se calcula precio − costo.'),
        ('Interpretación porcentajes', 'Participaciones y Score% se guardan como puntos porcentuales: 17,19 significa 17,19 %.'),
        *[(str(key), value) for key, value in parameters.items()],
    ], columns=['Parámetro', 'Valor'])
    frames = [('Costos', costs), ('Clasif ABC_XYZ', classification), ('Conclusiones', conclusions),
              ('Consolidado', consolidated), ('Pronosticos', models), ('Parametros', metadata)]
    money = {'MD', 'MOD', 'Costo fabricación', 'Precio venta', 'Utilidad unitaria', 'Utilidad total 52 semanas'}
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for name, frame in frames:
            frame.to_excel(writer, sheet_name=name, index=False, startrow=2)
            sheet = writer.sheets[name]
            sheet.cell(1, 1, f'MotoTrak · {name}')
            sheet.cell(1, 1).font = Font(size=16, bold=True, color='16324F')
            sheet.freeze_panes = 'A4'
            sheet.auto_filter.ref = f'A3:{get_column_letter(len(frame.columns))}{len(frame) + 3}'
            for cell in sheet[3]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='16324F')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.row_dimensions[3].height = 32
            for column_index, column in enumerate(frame.columns, 1):
                long_text = column in {'Texto', 'Valor', 'Fórmula Score%'}
                sheet.column_dimensions[get_column_letter(column_index)].width = 95 if long_text else min(42, max(18, len(str(column)) + 2))
                for row in sheet.iter_rows(min_row=4, min_col=column_index, max_col=column_index):
                    cell = row[0]
                    cell.alignment = Alignment(vertical='top', wrap_text=long_text)
                    if column in money:
                        cell.number_format = '"$" #,##0.00'
                    elif '%' in str(column) and column != 'Fórmula Score%':
                        cell.number_format = '0.00" %"'
                    elif column == 'Margen sobre venta':
                        cell.number_format = '0.00%'
                    elif 'RMSE' in str(column):
                        cell.number_format = '0.00'
            if 'Texto' in frame or 'Valor' in frame:
                for row_index in range(4, len(frame) + 4):
                    sheet.row_dimensions[row_index].height = 45
    return buffer.getvalue()
